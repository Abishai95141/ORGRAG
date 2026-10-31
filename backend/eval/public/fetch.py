"""Fetch the public-document eval set: real files from Microsoft's investor-relations
and sustainability pages, in the formats Strata parses. The files aren't committed
(they're Microsoft's); this script downloads them into eval/.data/public and checks
each against the hash it was written against, then derives the formats Microsoft
doesn't publish: a CSV from the financial statements, a scanned (image-only) PDF page,
and legacy .doc / .ppt copies (needs LibreOffice).

    .venv/bin/python -m eval.public.fetch
"""

import csv
import hashlib
import io
import shutil
import subprocess
import sys
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / ".data" / "public" / "microsoft" / "investor-relations"
CDN = "https://cdn-dynmedia-1.microsoft.com/is/content/microsoftcorp/"
CSR = CDN + "microsoft/msc/documents/presentations/CSR/"

# (file name, url, sha256 at the time the questions were written)
FILES: list[tuple[str, str, str]] = [
    (
        "fy26-q1-earnings-slides.pptx",
        CDN + "SlidesFY26Q1",
        "43ea442d4e881e1e07275c942357bca92b4f52b65336fa932a3254af334ecfb3",
    ),
    (
        "fy26-q1-earnings-call-transcript.docx",
        CDN + "TranscriptFY26Q1",
        "ecad78b9a38e33982b5c9ca285929921d4f82a64efc1f35cb93853311dd99280",
    ),
    (
        "fy26-q1-financial-statements.xlsx",
        CDN + "FinancialStatementFY26Q1",
        "a3448ba200e9f95ec7144b909e0c3e624d7ea492396f4e987eaa47c759798b5d",
    ),
    (
        "fy26-q2-outlook.pptx",
        CDN + "OutlookFY26Q2",
        "4cdeb54d3225460da03ef79fc43464ec6b1e8de4c8ee17456e5916b1502e5681",
    ),
    (
        "fy26-q1-press-release.docx",
        CDN + "PressReleaseFY26Q1",
        "4abab811f1e96eb56ebab7cc6b4448cf8c1e65255dcd114ec07825f545e040ac",
    ),
    (
        "fy26-q1-form-10-q.docx",
        CDN + "MSFT_FY26Q1_10Q",
        "fba0e3883b2f01dfe4f5a1a351688db9dbfe1d85168fe28706c6054efcb86592",
    ),
    (
        "2026-environmental-data-fact-sheet.pdf",
        CSR + "2026-Microsoft-Environmental-Data-Fact-Sheet-PDF.pdf",
        "0fd180837df00b3427c3eae484a492224b8f18c381e82a2aed98b04553a197f6",
    ),
    (
        "2026-environmental-sustainability-report.pdf",
        CSR + "2026-Microsoft-Environmental-Sustainability-Report-PDF.pdf",
        "20918e9fc4a02241f55d12f1f5b6228a832b706499f9963a1a3f91e3d3e693ba",
    ),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def download() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, url, expected in FILES:
        path = OUT / name
        if not path.exists():
            req = urllib.request.Request(url, headers={"User-Agent": "strata-eval"})
            with urllib.request.urlopen(req, timeout=120) as resp:
                path.write_bytes(resp.read())
        digest = sha256(path)
        status = "ok" if not expected or digest == expected else "CHANGED SINCE THE QUESTIONS"
        print(f"{status:<28} {name}  {digest[:16]}  {path.stat().st_size:,} bytes")


def derive() -> None:
    """The formats Microsoft doesn't publish, made from its own files."""
    import openpyxl
    import pypdfium2

    # CSV: the first sheet of the financial statements.
    book = openpyxl.load_workbook(OUT / "fy26-q1-financial-statements.xlsx", data_only=True)
    sheet = book.worksheets[0]
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in sheet.iter_rows(values_only=True):
        if any(v is not None and str(v).strip() for v in row):
            writer.writerow(["" if v is None else v for v in row])
    (OUT / "fy26-q1-income-statements.csv").write_text(buf.getvalue())

    # Scanned PDF: page 2 of the data fact sheet as an image only (no text layer).
    pdf = pypdfium2.PdfDocument(OUT / "2026-environmental-data-fact-sheet.pdf")
    image = pdf[1].render(scale=2.0).to_pil().convert("RGB")
    image.save(OUT / "2026-environmental-data-fact-sheet-page-2-scan.pdf", "PDF", resolution=144)

    # Legacy Office formats, through LibreOffice.
    soffice = shutil.which("soffice")
    if soffice is None:
        print("LibreOffice not found: skipping the .doc and .ppt copies", file=sys.stderr)
        return
    for src, fmt in (("fy26-q1-press-release.docx", "doc"), ("fy26-q2-outlook.pptx", "ppt")):
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                fmt,
                "--outdir",
                str(OUT / "legacy"),
                str(OUT / src),
            ],  # fmt: skip
            check=True,
            capture_output=True,
            timeout=180,
        )
    for f in (OUT / "legacy").iterdir():
        f.rename(OUT / f"{f.stem}-legacy{f.suffix}")
    (OUT / "legacy").rmdir()


if __name__ == "__main__":
    download()
    derive()
    for f in sorted(OUT.iterdir()):
        print(f"  {f.name}  {f.stat().st_size:,} bytes")
