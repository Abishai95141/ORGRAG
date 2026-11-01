import io
import json

import pytest

from app.ingestion.parsers import UnsupportedFileType, html_to_markdown, mime_for, parse


def test_mime_detection():
    assert mime_for("Report.PDF") == "application/pdf"
    assert mime_for("notes.markdown") == "text/markdown"
    with pytest.raises(UnsupportedFileType):
        mime_for("archive.zip")


def test_html_keeps_structure_and_drops_chrome():
    html = """<html><head><title>Docs</title><style>p{}</style></head><body>
    <nav>Menu</nav><main><h1>Install</h1><p>Run the installer.</p>
    <ul><li>Step one</li></ul><pre>make up</pre></main><footer>©</footer></body></html>"""
    md, title = html_to_markdown(html)
    assert title == "Docs"
    assert "# Install" in md and "- Step one" in md and "make up" in md
    assert "Menu" not in md and "©" not in md


def test_html_built_from_divs_keeps_its_text_once():
    """Most pages put text in <div>s and <span>s, not <p>s: all of it is kept, a table
    becomes a Markdown table, and a list item's paragraph isn't repeated."""
    html = """<body><div class="card"><span>Model</span><div>gpt-oss-20b</div></div>
    <section>Runtime: <b>llama.cpp</b> server<!-- a comment --></section>
    <table><tr><th>Plan</th><th>Seats</th></tr><tr><td>Growth</td><td>40</td></tr></table>
    <ul><li><p>Cut review time</p><ul><li>from days</li></ul></li></ul>
    <h5>Small heading</h5><button>Ask the twin</button></body>"""
    md, _ = html_to_markdown(html)
    blocks = md.split("\n\n")
    assert "Model" in blocks and "gpt-oss-20b" in blocks
    assert "Runtime: llama.cpp server" in blocks and "comment" not in md
    assert "| Plan | Seats |\n| --- | --- |\n| Growth | 40 |" in md
    assert "- Cut review time" in blocks and "- from days" in blocks
    assert md.count("Cut review time") == 1
    assert "##### Small heading" in blocks and "Ask the twin" in md


def test_csv_rows_become_records():
    md, _ = parse(b"plan,price\nGrowth,1200\nStarter,0\n", "text/csv")
    assert "### Row 1" in md and "**plan**: Growth" in md and "**price**: 0" in md


def test_json_is_pretty_printed():
    md, _ = parse(json.dumps({"a": [1, 2]}).encode(), "application/json")
    assert md.startswith("```json") and '"a"' in md


def test_markdown_title_detection():
    md, title = parse(b"# Security policy\n\nBody", "text/markdown")
    assert title == "Security policy" and "Body" in md


def test_docx_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Leave policy", level=1)
    d.add_paragraph("Twenty days per year.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Tenure", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "1-3", "20"
    buf = io.BytesIO()
    d.save(buf)
    md, title = parse(
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert title == "Leave policy"
    assert "## Leave policy" in md and "Twenty days" in md and "| Tenure | Days |" in md
