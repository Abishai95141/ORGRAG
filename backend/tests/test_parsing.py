"""Rich parsing: tables as passages, page numbers, weak-page signals, the optional
vision pass, and safety limits. Docling itself runs only with RUN_DOCLING=1 (loading
its models takes about a minute)."""

import io
import os
import zipfile

import pytest

from app.ingestion import docling_parser, vision
from app.ingestion.chunker import MAX_CHUNK_CHARS, chunk_markdown, page_marker, table_passages
from app.ingestion.docling_parser import ParseError, check_archive, debris_share


def test_large_tables_become_passages_that_repeat_the_header():
    rows = "\n".join(f"| R-{i:03d} | Item {i} | {i * 7} |" for i in range(120))
    table = f"| Code | Name | Count |\n|---|---|---|\n{rows}"
    pieces = table_passages(table, max_chars=1000)
    assert len(pieces) > 2
    assert all(p.startswith("| Code | Name | Count |\n|---|---|---|") for p in pieces)
    body = [line for p in pieces for line in p.splitlines()[2:]]
    assert body == rows.splitlines()  # every row once, in order

    md = f"# Inventory\n\nIntro.\n\n## Stock\n\n{table}\n\n## After\n\nClosing text."
    chunks = chunk_markdown(md, title="Inventory")
    stock = [c for c in chunks if c.heading == "Stock"]
    assert len(stock) > 1 and all(c.content.startswith("| Code |") for c in stock)
    assert chunks[-1].heading == "After" and chunks[-1].content == "Closing text."


def _padded(rows: list[list[str]]) -> str:
    """A table as Docling writes it: every cell padded to its column's widest."""
    widths = [max(len(r[i]) for r in rows) for i in range(len(rows[0]))]
    rule = "|" + "|".join("-" * (w + 2) for w in widths) + "|"

    def line(row: list[str]) -> str:
        return "| " + " | ".join(c.ljust(w) for c, w in zip(row, widths, strict=True)) + " |"

    return "\n".join([line(rows[0]), rule, *map(line, rows[1:])])


def test_padded_tables_fit_the_cap_and_keep_every_cell():
    guidance = " ".join(f"Sentence {i} of the guidance in one cell." for i in range(80))
    rows = [["Field", "Value"], ["Purpose", guidance]] + [
        [f"F{i}", f"value {i}"] for i in range(40)
    ]
    table = _padded(rows)
    assert len(table.splitlines()[0]) > 3000  # the padding alone would fill a passage
    chunks = chunk_markdown(f"# Record\n\n## Fields\n\n{table}", title="Record")
    head = "| Field | Value |\n|---|---|"
    assert all(len(c.content) <= MAX_CHUNK_CHARS for c in chunks)
    assert all(c.content.startswith(head + "\n") for c in chunks)  # padding gone, repeated
    body = "".join(c.content[len(head) :] for c in chunks)
    cells = "".join("|" + "|".join("".join(c.split()) for c in r) + "|" for r in rows[1:])
    assert "".join(body.split()) == cells  # every cell once, in order: split, never cut


def test_a_header_too_long_for_column_names_isnt_repeated():
    # A row Docling took for a header (a table continued from the page before).
    lead = " ".join(f"Point {i} carried over from the previous page." for i in range(30))
    table = f"| {lead} | Note |\n|---|---|\n" + "\n".join(
        f"| Row {i} {'detail ' * 10}| n{i} |" for i in range(60)
    )
    pieces = table_passages(table)
    assert len(pieces) > 1 and all(len(p) <= MAX_CHUNK_CHARS for p in pieces)
    assert pieces[0].startswith(f"| {lead} |") and not any(lead in p for p in pieces[1:])
    assert "\n".join(pieces) == table


def test_small_tables_stay_with_their_prose():
    md = "## Limits\n\nStandard keys:\n\n| Tier | Limit |\n|---|---|\n| Std | 1,000/h |"
    (chunk,) = chunk_markdown(md)
    assert "Standard keys:" in chunk.content and "| Std | 1,000/h |" in chunk.content


def test_chunks_record_the_page_they_start_on():
    md = (
        f"{page_marker(1)}\n\n# Report\n\nFirst page text.\n\n"
        f"{page_marker(2)}\n\n## Results\n\nSecond page text.\n\n"
        f"{page_marker(3)}\n\nThird page continues the results."
    )
    chunks = chunk_markdown(md, title="Report", target_words=4)
    assert [(c.heading, c.page) for c in chunks] == [
        ("", 1),
        ("Results", 2),
        ("Results", 3),
    ]
    assert all("<!--" not in c.content for c in chunks)  # markers aren't indexed


def test_ocr_debris_separates_garbled_text_from_clean_text():
    clean = (
        "Deloitte and Touche reviewed the greenhouse gas emissions reported in section one "
        "of this data fact sheet for the fiscal year ended June 30, 2025, and their review "
        "report is included at the end of the section together with the criteria used."
    )
    garbled = (
        "Ast snsg t a ng r ot o ng s f o the Environmental Data Fact Sheet are a "
        "compilation mm t t s ts t ti Initiative pden d o d o s in Section the single "
        "use pa dd t acd st as i Md o s licy is to begin tht i ese changes are included"
    )
    assert debris_share(clean) < 0.03
    assert debris_share(garbled) > docling_parser.WEAK_DEBRIS_SHARE
    assert debris_share("too few words to tell") == 0.0
    assert debris_share(clean.replace("gas", "mtCO2e")) < 0.03  # units aren't debris


def test_weak_pages_are_sparse_pictorial_or_garbled():
    weak = docling_parser._is_weak
    assert weak({"chars": 40, "pictures": 1, "picture_share": 0.1})  # a picture, little text
    assert weak({"chars": 4000, "pictures": 2, "picture_share": 0.55})  # mostly a chart
    assert weak({"chars": 6000, "pictures": 0, "picture_share": 0.0, "debris": 0.13})
    assert not weak({"chars": 40, "pictures": 0, "picture_share": 0.0})  # a short cover
    assert not weak({"chars": 3000, "pictures": 1, "picture_share": 0.05, "debris": 0.01})


def test_the_api_reads_the_stored_parse_of_older_and_pageless_documents():
    from app.api.serializers import document_parse

    assert document_parse({"parse": {"parser": "docling"}}) is None  # Word, HTML: no pages
    assert document_parse({}) is None and document_parse(None) is None
    old = document_parse({"parse": {"parser": "docling", "pages": 2}})  # before weak pages
    assert old is not None and old.weak_pages == [] and old.vision_pages is None


def _zip(entries: dict[str, bytes], compression=zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return buf.getvalue()


def test_office_files_that_expand_dangerously_are_refused():
    check_archive(_zip({"word/document.xml": b"<w:document/>" * 100}))  # ordinary
    with pytest.raises(ParseError, match="expands"):
        check_archive(_zip({"ppt/slides/slide1.xml": b"\0" * 50_000_000}))  # ratio bomb
    with pytest.raises(ParseError, match="valid Office"):
        check_archive(b"not a zip at all")


async def test_without_docling_rich_formats_are_refused_and_pdf_still_parses(monkeypatch):
    from app.ingestion import parsers

    monkeypatch.setattr(docling_parser, "available", lambda: False)
    monkeypatch.setattr(docling_parser, "legacy_office_available", lambda: False)
    assert ".xlsx" not in parsers.supported_extensions()
    with pytest.raises(parsers.UnsupportedFileType):
        parsers.mime_for("deck.pptx")
    with pytest.raises(parsers.UnsupportedFileType):
        await parsers.parse_document(b"x", parsers.RICH_EXTENSIONS[".xlsx"])
    parsed = await parsers.parse_document(b"# Title\n\nBody text.", "text/markdown")
    assert parsed.markdown.startswith("# Title") and parsed.title == "Title"


async def test_vision_replaces_only_the_weak_pages(monkeypatch):
    """Parent side: Docling's pages stay, weak pages take the transcript, cost is
    recorded, and a failed transcript keeps Docling's text."""

    def fake_convert(data, ext, render_weak):
        assert render_weak is True
        return {
            "title": "Deck",
            "pages": ["Clean first page.", "garbled o t s", "Chart page"],
            "numbers": [1, 2, 3],
            "weak": [2, 3],
            "renders": {2: b"png-2", 3: b"png-3"},
            "slides": True,
        }

    class InlinePool:
        def executor(self):
            return None

    calls = []

    async def fake_transcribe(png, model):
        calls.append(png)
        if png == b"png-3":
            raise RuntimeError("vision call failed")
        return "## Revenue\n\n| Segment | Revenue |\n|---|---|\n| Cloud | $30.9B |", 0.0004

    loop_calls = []

    async def run_inline(executor, fn, *args):
        loop_calls.append(fn)
        return fn(*args)

    import asyncio

    monkeypatch.setattr(docling_parser, "_convert", fake_convert)
    monkeypatch.setattr(docling_parser, "_pool", InlinePool())
    monkeypatch.setattr(vision, "enabled", lambda: True)
    monkeypatch.setattr(vision, "_transcribe", fake_transcribe)
    monkeypatch.setattr(asyncio.get_running_loop(), "run_in_executor", run_inline)

    parsed = await docling_parser.convert(b"%PDF-1.7", "application/pdf")
    assert sorted(calls) == [b"png-2", b"png-3"]
    assert parsed.info["vision_pages"] == [2] and parsed.info["vision_cost_usd"] == 0.0004
    assert parsed.info["weak_pages"] == [2, 3] and parsed.info["pages"] == 3
    assert "Clean first page." in parsed.markdown  # strong page kept
    assert "| Cloud | $30.9B |" in parsed.markdown and "garbled o t s" not in parsed.markdown
    assert "Chart page" in parsed.markdown  # failed transcript: Docling's text kept
    assert parsed.markdown.index(page_marker(2)) < parsed.markdown.index("| Cloud |")


def test_vision_is_off_without_the_setting_or_a_key(monkeypatch):
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "vision_parsing", False)
    assert vision.enabled() is False
    monkeypatch.setattr(get_settings(), "vision_parsing", True)
    monkeypatch.setattr(get_settings(), "anthropic_api_key", None)
    assert vision.enabled() is False


@pytest.mark.skipif(not os.environ.get("RUN_DOCLING"), reason="set RUN_DOCLING=1 to run Docling")
async def test_docling_keeps_word_tables_in_place_and_reads_excel():
    import docx
    import openpyxl

    if not docling_parser.available():
        pytest.skip("the `parsing` extra isn't installed")
    from app.ingestion.parsers import parse_document

    document = docx.Document()
    document.add_heading("Audit results", level=1)
    document.add_paragraph("Scores by area are below.")
    table = document.add_table(rows=3, cols=2)
    for r, (area, score) in enumerate([("Area", "Score"), ("Lab", "81%"), ("Exam rooms", "88%")]):
        table.cell(r, 0).text, table.cell(r, 1).text = area, score
    document.add_paragraph("Re-audit of the lab is on December 2.")
    buf = io.BytesIO()
    document.save(buf)
    word = await parse_document(buf.getvalue(), docling_parser.MIMES_BY_EXT[".docx"])
    md = word.markdown
    assert md.index("Scores by area") < md.index("| Lab") < md.index("Re-audit of the lab")

    book = openpyxl.Workbook()
    sheet = book.active
    for row in [("Cluster", "Nodes"), ("ingest-a", 36), ("query-b", 8)]:
        sheet.append(row)
    buf = io.BytesIO()
    book.save(buf)
    excel = await parse_document(buf.getvalue(), docling_parser.MIMES_BY_EXT[".xlsx"])
    assert "| ingest-a" in excel.markdown and "36" in excel.markdown
    docling_parser.shutdown()
