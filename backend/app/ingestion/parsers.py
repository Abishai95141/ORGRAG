"""Turn raw bytes into normalized Markdown. Each parser keeps headings so the
chunker can attach a heading path to every chunk.

`parse_document` is the entry point for indexing: PDF, Office files, email and images
go through Docling (docling_parser.py) when it's installed; everything else, and PDF
and DOCX without Docling, uses the light parsers here."""

import asyncio
import csv
import io
import json
import re
from pathlib import PurePosixPath
from typing import TYPE_CHECKING

from bs4 import BeautifulSoup, NavigableString, Tag
from bs4.element import PreformattedString

if TYPE_CHECKING:
    from app.ingestion.docling_parser import Parsed

SUPPORTED_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
    ".html": "text/html",
    ".htm": "text/html",
    ".csv": "text/csv",
    ".json": "application/json",
}
# Accepted only when Docling is installed (the `parsing` extra).
RICH_EXTENSIONS = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".eml": "message/rfc822",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}
# Legacy binary Office: also needs LibreOffice, which Docling converts them with.
LEGACY_EXTENSIONS = {
    ".doc": "application/msword",
    ".ppt": "application/vnd.ms-powerpoint",
    ".xls": "application/vnd.ms-excel",
}


class UnsupportedFileType(ValueError):
    pass


def supported_extensions() -> dict[str, str]:
    """Extension → mime for every type this deployment can parse."""
    from app.ingestion import docling_parser

    supported = dict(SUPPORTED_EXTENSIONS)
    if docling_parser.available():
        supported |= RICH_EXTENSIONS
    if docling_parser.legacy_office_available():
        supported |= LEGACY_EXTENSIONS
    return supported


def mime_for(filename: str) -> str:
    ext = PurePosixPath(filename.lower()).suffix
    supported = supported_extensions()
    if ext not in supported:
        raise UnsupportedFileType(f"Unsupported file type: {ext or filename}")
    return supported[ext]


async def parse_document(data: bytes, mime: str) -> "Parsed":
    """Markdown, detected title and parse details, using Docling for the formats it
    handles (and the light parsers otherwise)."""
    from app.ingestion import docling_parser

    if mime in docling_parser.DOCLING_TYPES and docling_parser.available():
        return await docling_parser.convert(data, mime)
    if mime in RICH_EXTENSIONS.values() or mime in LEGACY_EXTENSIONS.values():
        raise UnsupportedFileType(f"Parsing {mime} needs the `parsing` extra (Docling)")
    markdown, title = await asyncio.to_thread(parse, data, mime)  # CPU-bound
    return docling_parser.Parsed(markdown, title)


def parse(data: bytes, mime: str) -> tuple[str, str | None]:
    """Return (markdown, detected_title)."""
    if mime == "application/pdf":
        return _pdf(data), None
    if mime.endswith("wordprocessingml.document"):
        return _docx(data)
    if mime == "text/html":
        return html_to_markdown(data.decode("utf-8", errors="replace"))
    if mime == "text/csv":
        return _csv(data.decode("utf-8", errors="replace")), None
    if mime == "application/json":
        return _json(data.decode("utf-8", errors="replace")), None
    text = data.decode("utf-8", errors="replace")
    return text, _first_heading(text)


def _first_heading(md: str) -> str | None:
    m = re.search(r"^#{1,2}\s+(.+?)\s*$", md, re.MULTILINE)
    return m.group(1).strip() if m else None


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"## Page {i}\n\n{text}")
    return "\n\n".join(pages)


def _docx(data: bytes) -> tuple[str, str | None]:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    title = None
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style == "title":
            title = title or text
            lines.append(f"# {text}")
        elif style.startswith("heading"):
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 2
            lines.append(f"{'#' * min(level + 1, 6)} {text}")
            title = title or text
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            lines.append(_md_table(rows))
    return "\n\n".join(lines), title


# Elements that start a new block of text; anything else (span, a, strong, …) is inline
# and joins the text around it.
_BLOCK = frozenset(
    [
        "address", "article", "aside", "blockquote", "caption", "center", "dd", "details",
        "dialog", "div", "dl", "dt", "fieldset", "figcaption", "figure", "h1", "h2", "h3",
        "h4", "h5", "h6", "hr", "li", "main", "ol", "p", "pre", "section", "summary", "table",
        "tbody", "td", "tfoot", "th", "thead", "tr", "ul",
    ]
)  # fmt: skip
_CHROME = ["script", "style", "noscript", "svg", "nav", "footer", "header", "form", "template",
           "iframe", "canvas"]  # fmt: skip
_HEADING = re.compile(r"^h([1-6])$")


def html_to_markdown(html: str) -> tuple[str, str | None]:
    """Markdown from HTML: headings, lists, code, quotes and tables keep their shape, and
    every other block's text is a paragraph. (Keeping only <p>, <li> and headings lost
    pages built from <div>s and <span>s, which is most of them.)"""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(_CHROME):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else None
    root = soup.find("main") or soup.find("article") or soup.body or soup
    blocks: list[str] = []
    _blocks(root, blocks)
    deduped: list[str] = []
    for block in blocks:
        if not deduped or deduped[-1] != block:
            deduped.append(block)
    return "\n\n".join(deduped), title


def _text(el) -> str:
    return " ".join(el.get_text(" ").split())


def _has_block(el: Tag) -> bool:
    return any(d.name in _BLOCK for d in el.find_all(True))


def _blocks(el: Tag, out: list[str]) -> None:
    inline: list[str] = []

    def flush() -> None:
        text = " ".join(" ".join(inline).split())
        if text:
            out.append(text)
        inline.clear()

    for child in el.children:
        if isinstance(child, PreformattedString):  # comments, doctype, CDATA
            continue
        if isinstance(child, NavigableString):
            inline.append(str(child))
            continue
        if not isinstance(child, Tag):
            continue
        name = child.name
        if name not in _BLOCK and not _has_block(child):
            inline.append(child.get_text(" "))
            continue
        flush()
        if heading := _HEADING.match(name):
            if text := _text(child):
                out.append(f"{'#' * int(heading.group(1))} {text}")
        elif name == "pre":
            out.append(f"```\n{child.get_text().strip(chr(10))}\n```")
        elif name == "li":
            if not _has_block(child):
                if text := _text(child):
                    out.append(f"- {text}")
            else:
                nested: list[str] = []
                _blocks(child, nested)
                if nested:
                    out.append(f"- {nested[0]}")
                    out.extend(nested[1:])
        elif name == "blockquote":
            quoted: list[str] = []
            _blocks(child, quoted)
            out.extend(f"> {q}" for q in quoted)
        elif name == "table":
            rows = [[_text(c) for c in tr.find_all(["th", "td"])] for tr in child.find_all("tr")]
            rows = [r for r in rows if any(r)]
            if len(rows) > 1 and max(len(r) for r in rows) > 1:
                out.append(_md_table(rows))
            else:
                out.extend(" ".join(r) for r in rows)
        elif name != "hr":
            _blocks(child, out)
    flush()


def _md_table(rows: list[list[str]]) -> str:
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([head, sep, *body])


def _csv(text: str) -> str:
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    # Render each record as a key/value block so rows chunk and retrieve independently.
    blocks = []
    for i, row in enumerate(body, start=1):
        pairs = [f"- **{h}**: {v}" for h, v in zip(header, row, strict=False) if v.strip()]
        blocks.append(f"### Row {i}\n" + "\n".join(pairs))
    return "\n\n".join(blocks) if blocks else _md_table(rows)


def _json(text: str) -> str:
    try:
        return "```json\n" + json.dumps(json.loads(text), indent=2)[:200_000] + "\n```"
    except json.JSONDecodeError:
        return text
